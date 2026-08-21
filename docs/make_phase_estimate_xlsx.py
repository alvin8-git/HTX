#!/usr/bin/env python3
"""Build docs/HTX_phase1_phase2_estimate.xlsx.

Three columns only - item, description, hours - and written for a reader who will never open the
code. Edit the tables below and re-run. Sub-item detail lives in docs/phase1_triage_scope.md;
rates and money live in docs/effort_and_cost_estimate.md. ponytail: no config file, the data is
the file.
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

BLUE, LIGHT = '00539B', 'EDF2F8'
H = Font(color='FFFFFF', bold=True, size=11)
B = Font(bold=True)
TITLE = Font(bold=True, size=14, color='00539B')
FILL_H = PatternFill('solid', fgColor=BLUE)
FILL_SUB = PatternFill('solid', fgColor=LIGHT)
WRAP = Alignment(wrap_text=True, vertical='top')
TOP = Alignment(vertical='top', horizontal='center')
RULE = Border(bottom=Side('thin', color='BBBBBB'))

# (item, description, low, high)
P1_BUILT = [
 ('Rules and reference knowledge base',
  '447 curated entries: which organisms are a threat, which resistance genes matter, and which '
  'markers separate a dangerous strain from a harmless relative. Every entry is a judgement made '
  'from the published literature, not data entry. This is what the system knows.', 95, 130),
 ('The analysis engine',
  'The script itself. It reads the sequencer’s output, works out what is in the sample, and puts '
  'every organism through a fixed series of checks to reach one of four verdicts — no action, '
  'monitor, investigate, escalate. It also records, in plain English, which rule produced each '
  'verdict, and compares each new swab against what is normal for that site.', 165, 225),
 ('The analyst’s report',
  'A self-contained report per swab: the flagged organisms first, then the evidence behind each '
  'one. Written so an analyst can re-trace any verdict without help. A matching data file is '
  'produced at the same time for the future database.', 85, 120),
 ('Testing and fault-finding',
  '14 real samples across three independent datasets — field swabs, a control mix with a known '
  'answer, and a long-read stool run. Ten genuine faults were found and fixed. Built-in checks now '
  'run on demand so a change cannot silently break a verdict.', 105, 155),
 ('Supporting tools and packaging',
  'Ten smaller tools around the main script (control checking, briefing slides, rule-base upkeep), '
  'plus the packaging that lets the same script run on a laptop, a server or the sequencing '
  'appliance and give the same answer.', 75, 120),
 ('Documentation',
  'Around 120 pages: how the system decides, what it can and cannot see, why each rule exists, and '
  'the record of every review and fix.', 65, 100),
]
P1_TODO = [
 ('Complete the rule base',
  'Two gaps. The strain-level markers cover 1 of 25 watchlist organisms; the other 24 need a '
  'literature pass. And the collection date, site and sequencing run ID are still blank on the '
  'sample sheet — a few hours of work that must happen before the next samples are collected, '
  'because it cannot be recovered afterwards and it blocks half of Phase 2.', 44, 78),
 ('Detect engineered organisms',
  'Look for combinations of genes that do not occur together in nature, and measure how often that '
  'raises a false alarm. The cheapest item on either phase, and the only one that addresses '
  'deliberate modification rather than natural occurrence.', 20, 40),
 ('Formal validation and sign-off',
  'A written, repeatable acceptance test with pass/fail thresholds — the evidence that the system '
  'works, rather than the belief that it does.', 60, 100),
 ('Production readiness and handover',
  'Sensible behaviour on bad input, clear errors and logs so a schedule can run it instead of a '
  'person, plus training material for the analysts who will read the reports.', 45, 75),
]
P2 = [
 ('RNA virus coverage',
  'The assay currently reads DNA only, so 25 of 46 priority agents — including most of the viruses '
  '— are invisible to it. This is the single largest gain in capability in the whole programme. '
  'It also depends on laboratory work that is not costed here and may exceed the analysis effort.',
  120, 200),
 ('Central database and automatic processing',
  'Every analysed sample, verdict and piece of evidence stored in one place, with a full record of '
  'which version of the rules produced it. New sequencing runs are picked up and analysed without '
  'anyone typing a command, and jobs are queued so a busy day does not overwhelm the machine.',
  260, 380),
 ('Managing dashboard',
  'The screen an analyst opens each morning: every sample processed, the verdict breakdown, the '
  'sites being watched, and a click-through into any individual report.', 180, 270),
 ('Tracking sites over time, and forecasting',
  'Same site, repeated visits: is this organism new, is it growing, or has it always been there? '
  'Includes correcting for the fact that separate sequencing runs are not directly comparable, and '
  'turning the history into a forward signal that flags a reading outside the expected range.',
  180, 280),
 ('Security, hosting and reliability',
  'Who may see what, a tamper-evident record of who looked at or changed anything, the system '
  'watching its own health, and standing it up, backing it up and proving it can be restored.',
  200, 320),
 ('Validation, documentation and training',
  'Acceptance testing of the whole platform rather than just the script, plus operator and analyst '
  'material. The long pole, and the line most often cut — cutting it moves the cost to the first '
  'incident.', 230, 370),
]

def sheet(wb, name, title, subtitle, groups):
    ws = wb.create_sheet(name)
    ws['A1'] = title; ws['A1'].font = TITLE
    ws['A2'] = subtitle; ws['A2'].alignment = WRAP
    ws.merge_cells('A2:C2'); ws.row_dimensions[2].height = 45
    for i, h in enumerate(['Item', 'Description', 'Hours'], 1):
        c = ws.cell(4, i, h); c.font = H; c.fill = FILL_H; c.alignment = WRAP
    r, first, spans = 5, 5, []
    for grp, items in groups:
        if grp:
            c = ws.cell(r, 1, grp); c.font = B; c.fill = FILL_SUB
            ws.cell(r, 2).fill = FILL_SUB; ws.cell(r, 3).fill = FILL_SUB
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
            r += 1
        gfirst = r
        for item, desc, lo, hi in items:
            ws.cell(r, 1, item).alignment = WRAP
            ws.cell(r, 2, desc).alignment = WRAP
            ws.cell(r, 3, f'{lo}–{hi}').alignment = TOP
            ws.cell(r, 4, lo); ws.cell(r, 5, hi)          # hidden, drives the totals
            for i in range(1, 4):
                ws.cell(r, i).border = RULE
            ws.row_dimensions[r].height = 62
            r += 1
        spans.append((gfirst, r - 1))
        if grp:
            ws.cell(r, 1, 'Subtotal').font = B
            ws.cell(r, 3, f'=D{r}&"–"&E{r}').font = B
            ws.cell(r, 4, f'=SUM(D{gfirst}:D{r-1})')
            ws.cell(r, 5, f'=SUM(E{gfirst}:E{r-1})')
            for i in range(1, 4):
                ws.cell(r, i).fill = FILL_SUB
            r += 2
    ws.cell(r, 1, 'TOTAL').font = B
    ws.cell(r, 3, f'=D{r}&"–"&E{r}').font = B
    ws.cell(r, 4, '=' + '+'.join(f'SUM(D{a}:D{b})' for a, b in spans))
    ws.cell(r, 5, '=' + '+'.join(f'SUM(E{a}:E{b})' for a, b in spans))
    for i in range(1, 4):
        ws.cell(r, i).fill = FILL_SUB
    for col, w in zip('ABC', (34, 88, 12)):
        ws.column_dimensions[col].width = w
    ws.column_dimensions['D'].hidden = True
    ws.column_dimensions['E'].hidden = True
    ws.freeze_panes = 'A5'
    return spans

wb = Workbook(); wb.remove(wb.active)
P1SPANS = sheet(wb, 'Phase 1',
    'Phase 1 — the analysis script, current DNA dataset',
    'The script that turns one sequencing run into one interpretable report, the rule base it '
    'judges against, and the testing behind it — for the DNA assay as it stands today. RNA, the '
    'dashboard, the database and any over-time analysis are Phase 2. Hours are the effort to build '
    'this conventionally, not the wall-clock time already spent.',
    [('ALREADY BUILT', P1_BUILT), ('REMAINING TO CLOSE PHASE 1', P1_TODO)])
P2SPANS = sheet(wb, 'Phase 2',
    'Phase 2 — hosted platform, RNA, tracking over time, forecasting',
    'Everything that turns the script into a running surveillance service, plus the RNA coverage '
    'the current assay cannot provide. All of it is new build.',
    [(None, P2)])

sm = wb.create_sheet('Summary', 0)
sm['A1'] = 'HTX biosurveillance — effort summary'; sm['A1'].font = TITLE
for i, h in enumerate(['Phase', 'Hours', 'Mid-point'], 1):
    c = sm.cell(3, i, h); c.font = H; c.fill = FILL_H
# ranges come from where the rows actually landed - hardcoding them is how a row goes missing
srows = [('Phase 1 — already built', "'Phase 1'!") + P1SPANS[0],
         ('Phase 1 — remaining', "'Phase 1'!") + P1SPANS[1],
         ('Phase 2 — platform, RNA, tracking, forecasting', "'Phase 2'!") + P2SPANS[0]]
r = 4
for label, sh, a, b in srows:
    sm.cell(r, 1, label)
    sm.cell(r, 4, f'=SUM({sh}D{a}:D{b})'); sm.cell(r, 5, f'=SUM({sh}E{a}:E{b})')
    sm.cell(r, 2, f'=D{r}&"–"&E{r}')
    sm.cell(r, 3, f'=ROUND(AVERAGE(D{r}:E{r}),0)')
    r += 1
sm.cell(r, 1, 'TOTAL, both phases').font = B
sm.cell(r, 4, f'=SUM(D4:D{r-1})'); sm.cell(r, 5, f'=SUM(E4:E{r-1})')
sm.cell(r, 2, f'=D{r}&"–"&E{r}').font = B
sm.cell(r, 3, f'=ROUND(AVERAGE(D{r}:E{r}),0)').font = B
for i in range(1, 4):
    sm.cell(r, i).fill = FILL_SUB
notes = ['', 'Not costed anywhere in this workbook',
 '   • Laboratory development for RNA coverage: reagents, protocol work and instrument time. '
 'This may exceed the analysis effort attached to it and is the largest single risk to the estimate.',
 '   • Sequencing consumables and instrument time for revalidation runs.',
 '   • Hardware. Compute is modest and storage is trivial.',
 '   • Accreditation or regulatory submission, if it is ever required. That is its own programme.',
 '', 'How to read the "already built" figure',
 '   It is replacement value, not spend. The work was done with heavy AI assistance, so the actual '
 'time taken was a fraction of it. The figure answers "what would this have cost to build "'
 '"conventionally" — the right number for valuing the asset or for a build-versus-buy comparison. '
 'It is not a timesheet.',
 '', 'Two things that will move these numbers',
 '   1. Filling in the blank collection date, site and run ID costs under a day and unblocks '
 'roughly half of Phase 2. It has to be done before the next samples are collected.',
 '   2. About half of Phase 2 is not scientific work. Costing all of it at a senior scientist rate '
 'both overstates the budget and understates the risk.',
 '', 'Item-by-item detail: docs/phase1_triage_scope.md. Rates and money: '
 'docs/effort_and_cost_estimate.md.']
r += 2
for line in notes:
    c = sm.cell(r, 1, line); c.alignment = WRAP
    if line and not line.startswith(' '):
        c.font = B
    elif line:
        sm.row_dimensions[r].height = 32
    sm.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    r += 1
for col, w in zip('ABC', (52, 16, 12)):
    sm.column_dimensions[col].width = w
sm.column_dimensions['D'].hidden = True
sm.column_dimensions['E'].hidden = True

out = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'HTX_phase1_phase2_estimate.xlsx')
wb.save(out)
print('wrote', out, P1SPANS, P2SPANS)
